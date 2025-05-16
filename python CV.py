from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Marges aanpassen
section = doc.sections[0]
section.top_margin = Inches(0.5)
section.bottom_margin = Inches(0.5)
section.left_margin = Inches(0.5)
section.right_margin = Inches(0.5)

# Titelblok
title = doc.add_heading('SEVBAN BURGAZ', level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].font.color.rgb = RGBColor(0, 102, 204)
title.runs[0].font.size = Pt(24)

subtitle = doc.add_paragraph('SOFTWARE DEVELOPER')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(14)
subtitle.runs[0].font.color.rgb = RGBColor(100, 100, 100)

doc.add_paragraph('\n')

# CV in twee kolommen (tabel)
table = doc.add_table(rows=1, cols=2)
table.autofit = False
table.columns[0].width = Inches(2.5)
table.columns[1].width = Inches(4.5)

# Linker kolom
left = table.cell(0, 0)
left.text = ''
left_para = left.paragraphs[0]
left_para.add_run('PERSOONLIJK PROFIEL\n').bold = True
left.add_paragraph("Ik ben een professionele en meedenkende persoon die sterk is in multitasken en het efficiënt uitvoeren van opdrachten. Ik werk doelgericht, ben gefocust en weet van aanpakken.")

left.add_paragraph('\nCONTACT\n').runs[0].bold = True
left.add_paragraph("Koningspleinhof 20\n+31 06 15364805\nsevban_38@outlook.com\nLinkedIn: Sevban Burgaz")

left.add_paragraph('\nOPLEIDING\n').runs[0].bold = True
left.add_paragraph("Da Vinci College\nSoftware Developer – heden")
left.add_paragraph("Insula College – Koningstraat\nVmbo-TL/GL")

# Rechter kolom
right = table.cell(0, 1)
right.text = ''
right.paragraphs[0].add_run('VAARDIGHEDEN\n').bold = True
right.add_paragraph("- Multitasking\n- Professionele werkhouding\n- Gericht en gefocust werken")

right.add_paragraph('\nTECHNISCHE VAARDIGHEDEN\n').runs[0].bold = True
right.add_paragraph("Python, CSS, HTML, JavaScript\nSQLite, MySQL")

right.add_paragraph('\nTALEN\n').runs[0].bold = True
right.add_paragraph("Turks – 60%\nEngels – 70%\nNederlands – 70%")

right.add_paragraph('\nWERKERVARING\n').runs[0].bold = True
right.add_paragraph("Albert Heijn – Achterom\nVakkenvuller (9 maanden)")
right.add_paragraph("- Schappen vullen en houdbaarheid controleren\n- Klanten helpen")

right.add_paragraph("New York Pizza – Transvaalstraat\nBezorger / Pizzamaker (6 maanden)")
right.add_paragraph("- Pizza’s bereiden en bezorgen\n- Orders verwerken")

# Placeholder voor foto
doc.add_paragraph('\n[Voeg hier een profielfoto toe]', style='Intense Quote')

# Opslaan
doc.save("Sevban_Burgaz_CV.docx")
